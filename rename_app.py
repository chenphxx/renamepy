# -*- coding: utf-8 -*-
"""主窗口：批量重命名工具的整体界面、预览与处理流程。"""

import csv
import os
import shutil
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from config import COLORS, DARK, LIGHT, FONT_FAMILY, WINDOW_MIN_SIZE, WINDOW_SIZE, WINDOW_TITLE
from file_ops import (
    clear_files_in_folder,
    open_folder,
    prepare_inplace_temps,
    restore_inplace_temp,
    restore_inplace_temps,
)
from folder_block import FolderBlock

# 预览支持的文件类型
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".ico", ".tif", ".tiff"}
TEXT_EXTS = {".txt", ".md", ".log", ".py", ".json", ".xml", ".yml", ".yaml", ".ini", ".cfg", ".conf", ".html", ".css", ".js", ".srt"}
TABLE_EXTS = {".csv", ".tsv", ".xlsx"}
MONO_FONT = ("Consolas", 10) if os.name == "nt" else ("Courier", 10)
MAX_PREVIEW_ROWS = 100
MAX_PREVIEW_BYTES = 200 * 1024
PREVIEW_HEIGHT = 210
PREVIEW_MIN_HEIGHT = 120

# 会被逐个重刷颜色的 tk 控件所用到的主题键（ttk 样式由 _setup_style 统一更新）
WALKED_COLOR_KEYS = (
    "bg", "card", "card_border", "header", "header_border",
    "text", "text_secondary", "accent", "preview_bg", "menu_bg", "menu_hover",
)


class RenameApp:
    def __init__(self, root):
        self.root = root
        root.title(WINDOW_TITLE)
        root.geometry(WINDOW_SIZE)
        root.minsize(*WINDOW_MIN_SIZE)
        root.configure(bg=COLORS["bg"])

        self.folder_blocks = []
        self._next_folder_id = 1
        self.global_action = None  # None / "overwrite" / "skip"
        self.cancel_flag = False
        self.theme = "dark"

        self._preview_path = None
        self._preview_img = None
        self._preview_rendered_size = None
        self._drag_start_y = None
        self._drag_start_h = None

        self._setup_style()
        self._build_ui()

    # ---------------- 样式（Discord 风格） ----------------
    def _setup_style(self):
        self.style = ttk.Style()
        try:
            self.style.theme_use("clam")
        except Exception:
            pass

        c = COLORS
        font = (FONT_FAMILY, 10)

        # 下拉框弹出列表配色
        self.root.option_add("*TCombobox*Listbox.background", c["menu_bg"])
        self.root.option_add("*TCombobox*Listbox.foreground", c["text"])
        self.root.option_add("*TCombobox*Listbox.selectBackground", c["accent"])
        self.root.option_add("*TCombobox*Listbox.selectForeground", "#FFFFFF")

        self.style.configure(".", font=font, background=c["bg"], foreground=c["text"])
        self.style.configure("TFrame", background=c["bg"])
        self.style.configure("TLabel", background=c["bg"], foreground=c["text"])
        self.style.configure("Muted.TLabel", background=c["bg"], foreground=c["text_secondary"])
        self.style.configure("Card.TLabel", background=c["card"], foreground=c["text"])
        self.style.configure("CardMuted.TLabel", background=c["card"], foreground=c["text_secondary"])

        # 主按钮（Discord 紫蓝）
        self.style.configure(
            "Primary.TButton",
            background=c["accent"], foreground="#FFFFFF",
            bordercolor=c["accent"], lightcolor=c["accent"], darkcolor=c["accent"],
            relief="flat", padding=(18, 8), font=(FONT_FAMILY, 10, "bold"),
        )
        self.style.map("Primary.TButton", background=[("pressed", c["accent_hover"]), ("active", c["accent_hover"])])

        # 次要按钮（Discord 灰）
        self.style.configure(
            "Ghost.TButton",
            background=c["button_secondary"], foreground=c["button_text"],
            bordercolor=c["button_secondary"], lightcolor=c["button_secondary"], darkcolor=c["button_secondary"],
            relief="flat", padding=(12, 7),
        )
        self.style.map(
            "Ghost.TButton",
            background=[("pressed", c["button_secondary_hover"]), ("active", c["button_secondary_hover"])],
        )

        # 危险按钮（删除，红字）
        self.style.configure(
            "Danger.TButton",
            background=c["card"], foreground=c["danger"],
            bordercolor=c["card_border"], lightcolor=c["card_border"], darkcolor=c["card_border"],
            relief="flat", padding=(10, 6),
        )
        self.style.map("Danger.TButton", background=[("pressed", c["danger_hover"]), ("active", c["danger_hover"])])

        # 输入框
        self.style.configure(
            "TEntry",
            fieldbackground=c["input"], foreground=c["text"],
            bordercolor=c["input_border"], lightcolor=c["input_border"], darkcolor=c["input_border"],
            relief="flat", padding=6,
        )
        self.style.map(
            "TEntry",
            bordercolor=[("focus", c["accent"])],
            lightcolor=[("focus", c["accent"])],
            darkcolor=[("focus", c["accent"])],
        )

        # 下拉框
        self.style.configure(
            "TCombobox",
            fieldbackground=c["input"], background=c["input"], foreground=c["text"],
            bordercolor=c["input_border"], lightcolor=c["input_border"], darkcolor=c["input_border"],
            arrowcolor=c["text_secondary"], relief="flat", padding=4,
        )
        self.style.map(
            "TCombobox",
            fieldbackground=[("readonly", c["input"])],
            foreground=[("readonly", c["text"])],
        )

        # 文件列表
        self.style.configure(
            "Treeview",
            background=c["card"], fieldbackground=c["card"], foreground=c["text"],
            bordercolor=c["card_border"], relief="flat", rowheight=26,
        )
        self.style.configure(
            "Treeview.Heading",
            background=c["heading_bg"], foreground=c["text_secondary"],
            relief="flat", font=(FONT_FAMILY, 10, "bold"),
        )
        self.style.map(
            "Treeview",
            background=[("selected", c["accent_selected"])],
            foreground=[("selected", c["accent_selected_fg"])],
        )
        self.style.map("Treeview.Heading", background=[("active", c["heading_hover"])])

        # 进度条 / 滚动条
        self.style.configure(
            "Horizontal.TProgressbar",
            troughcolor=c["progress_trough"], background=c["progress_bar"],
            bordercolor=c["progress_trough"], lightcolor=c["progress_bar"], darkcolor=c["progress_bar"],
            thickness=8,
        )
        self.style.configure(
            "Vertical.TScrollbar",
            troughcolor=c["bg"], background=c["scrollbar"],
            bordercolor=c["bg"], arrowcolor=c["text_secondary"],
        )
        self.style.configure(
            "Horizontal.TScrollbar",
            troughcolor=c["bg"], background=c["scrollbar"],
            bordercolor=c["bg"], arrowcolor=c["text_secondary"],
        )

    # ---------------- UI 布局 ----------------
    def _build_ui(self):
        root = self.root
        c = COLORS

        # ---- 顶部标题区 ----
        header = tk.Frame(root, bg=c["header"], highlightbackground=c["header_border"], highlightthickness=1)
        header.pack(side="top", fill="x")
        tk.Label(
            header, text=WINDOW_TITLE, bg=c["header"], fg=c["text"],
            font=(FONT_FAMILY, 15, "bold"),
        ).pack(side="left", padx=(20, 14), pady=12)
        tk.Label(
            header, text="按「前缀 + 序号」批量重命名文件", bg=c["header"], fg=c["text_secondary"],
            font=(FONT_FAMILY, 10),
        ).pack(side="left")

        self.btn_theme = ttk.Button(header, text="浅色模式", style="Ghost.TButton", command=self.toggle_theme)
        self.btn_theme.pack(side="right", padx=(0, 16))

        # ---- 顶部控制区 ----
        top_frame = ttk.Frame(root, padding=(16, 12, 16, 4))
        top_frame.pack(side="top", fill="x")

        self.btn_add_folder = ttk.Button(top_frame, text="+ 添加文件夹", style="Ghost.TButton", command=self.add_folder)
        self.btn_add_folder.pack(side="right")

        ttk.Label(top_frame, text="文件名前缀").pack(side="left")
        self.entry_prefix = ttk.Entry(top_frame, width=20)
        self.entry_prefix.pack(side="left", padx=(8, 4))

        ttk.Label(top_frame, text="起始数字").pack(side="left", padx=(14, 4))
        self.entry_start = ttk.Entry(top_frame, width=8)
        self.entry_start.insert(0, "1")
        self.entry_start.pack(side="left")

        self.btn_rename = ttk.Button(top_frame, text="开始处理", style="Primary.TButton", command=self.rename_and_save)
        self.btn_rename.pack(side="left", padx=(20, 8))

        self.btn_cancel = ttk.Button(top_frame, text="取消", style="Ghost.TButton", command=self.cancel_process)
        self.btn_cancel.pack(side="left")

        # ---- 提示行 ----
        hint_frame = ttk.Frame(root, padding=(16, 2, 16, 6))
        hint_frame.pack(side="top", fill="x")
        ttk.Label(
            hint_frame,
            text="提示：保存位置选择与原文件夹相同的目录，即可直接在原文件夹内完成重命名",
            style="Muted.TLabel",
        ).pack(anchor="w")

        # ---- 中间可滚动区域（文件夹卡片） ----
        container = tk.Frame(root, bg=c["bg"])
        container.pack(side="top", fill="both", expand=True, padx=14, pady=(0, 4))

        self.canvas = tk.Canvas(container, bg=c["bg"], highlightthickness=0)
        self.v_scroll = ttk.Scrollbar(container, orient="vertical", command=self.canvas.yview)
        self.canvas.configure(yscrollcommand=self.v_scroll.set)
        self.canvas.pack(side="left", fill="both", expand=True)
        self.v_scroll.pack(side="right", fill="y")

        self.blocks_frame = ttk.Frame(self.canvas)
        self.canvas_window = self.canvas.create_window((0, 0), window=self.blocks_frame, anchor="nw")
        self.blocks_frame.bind("<Configure>", lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        self.canvas.bind("<Configure>", lambda e: self.canvas.itemconfig(self.canvas_window, width=e.width))

        # 鼠标滚轮绑定（仅内容高于 canvas 时滚动）
        self._bind_mousewheel(self.canvas)

        # ---- 底部状态区（先 pack，位于最底部） ----
        bottom = ttk.Frame(root, padding=(16, 8, 16, 10))
        bottom.pack(side="bottom", fill="x")

        self.progress = ttk.Progressbar(bottom, orient="horizontal", mode="determinate")
        self.progress.pack(side="left", fill="x", expand=True, padx=(0, 12))

        self.lbl_percent = ttk.Label(bottom, text="0%", width=6)
        self.lbl_percent.pack(side="left", padx=(0, 12))

        self.lbl_count = ttk.Label(bottom, text="已处理 0 / 0", width=18)
        self.lbl_count.pack(side="left")

        # ---- 文件预览区（位于状态栏上方） ----
        preview_outer = tk.Frame(root, bg=c["bg"])
        preview_outer.pack(side="bottom", fill="x", padx=14, pady=(0, 8))

        self.preview_frame = tk.Frame(
            preview_outer, bg=c["card"],
            highlightbackground=c["card_border"], highlightthickness=1,
        )
        self.preview_frame.configure(height=PREVIEW_HEIGHT)
        self.preview_frame.pack_propagate(False)
        self.preview_frame.pack(fill="x")

        preview_header = tk.Frame(self.preview_frame, bg=c["card"])
        preview_header.pack(fill="x", padx=14, pady=(8, 6))
        self.lbl_preview_title = tk.Label(
            preview_header, text="文件预览", bg=c["card"], fg=c["text"],
            font=(FONT_FAMILY, 11, "bold"),
        )
        self.lbl_preview_title.pack(side="left")
        self.lbl_preview_info = tk.Label(
            preview_header, text="点击左侧文件列表中的文件查看预览", bg=c["card"], fg=c["text_secondary"],
        )
        self.lbl_preview_info.pack(side="left", padx=(14, 0))

        self.preview_body = tk.Frame(self.preview_frame, bg=c["card"])
        self.preview_body.pack(fill="both", expand=True, pady=(0, 10))
        self.preview_body.bind("<Configure>", self._on_preview_resize)

        # ---- 分隔条（拖动调整预览区高度） ----
        self.divider = tk.Frame(root, bg=c["card_border"], cursor="sb_v_double_arrow", height=6)
        self.divider.pack(side="bottom", fill="x", padx=14)
        self.divider.bind("<Button-1>", self._divider_press)
        self.divider.bind("<B1-Motion>", self._divider_drag)
        self.divider.bind("<Enter>", lambda e: self.divider.configure(bg=c["accent"]))
        self.divider.bind("<Leave>", lambda e: self.divider.configure(bg=c["card_border"]))

        # 初始化时添加一个文件夹卡片，并显示预览占位
        self.add_folder()
        self.show_preview(None)

    # ---------------- 主题切换 ----------------
    def toggle_theme(self):
        old = dict(COLORS)
        if self.theme == "dark":
            new = LIGHT
            self.theme = "light"
        else:
            new = DARK
            self.theme = "dark"

        # 旧颜色 -> 新颜色 映射，仅包含 tk 控件实际用到的主题键
        mapping = {v: new[k] for k, v in old.items() if k in WALKED_COLOR_KEYS and v != new[k]}
        COLORS.clear()
        COLORS.update(new)

        self._setup_style()
        self._apply_colors_to_widgets(mapping)
        self._refresh_tree_tags()
        self.btn_theme.config(text="浅色模式" if self.theme == "dark" else "深色模式")
        if self._preview_path:
            self.show_preview(self._preview_path)

    def _apply_colors_to_widgets(self, mapping):
        def recolor(widget):
            try:
                if isinstance(widget, tk.Label):
                    for opt in ("bg", "fg"):
                        v = str(widget.cget(opt))
                        if v in mapping:
                            widget.configure(**{opt: mapping[v]})
                elif isinstance(widget, tk.Frame):
                    for opt in ("bg", "highlightbackground"):
                        v = str(widget.cget(opt))
                        if v in mapping:
                            widget.configure(**{opt: mapping[v]})
                elif isinstance(widget, tk.Canvas):
                    v = str(widget.cget("bg"))
                    if v in mapping:
                        widget.configure(bg=mapping[v])
                elif isinstance(widget, tk.Text):
                    for opt in ("bg", "fg", "insertbackground"):
                        v = str(widget.cget(opt))
                        if v in mapping:
                            widget.configure(**{opt: mapping[v]})
                elif isinstance(widget, tk.Menu):
                    for opt in ("bg", "fg", "activebackground", "activeforeground"):
                        v = str(widget.cget(opt))
                        if v in mapping:
                            widget.configure(**{opt: mapping[v]})
            except tk.TclError:
                pass
            for child in widget.winfo_children():
                recolor(child)

        recolor(self.root)
        self.root.configure(bg=COLORS["bg"])

    def _refresh_tree_tags(self):
        for block in self.folder_blocks:
            block.tree.tag_configure("oddrow", background=COLORS["card"])
            block.tree.tag_configure("evenrow", background=COLORS["row_alt"])

    # ---------------- 预览区拖拽 ----------------
    def _divider_press(self, event):
        self._drag_start_y = event.y_root
        self._drag_start_h = self.preview_frame.winfo_height()

    def _divider_drag(self, event):
        if self._drag_start_y is None:
            return
        delta = self._drag_start_y - event.y_root
        min_h = PREVIEW_MIN_HEIGHT
        max_h = max(min_h, self.root.winfo_height() - 420)
        new_h = max(min_h, min(self._drag_start_h + delta, max_h))
        self.preview_frame.configure(height=new_h)
        self.root.update_idletasks()

    def _on_preview_resize(self, event):
        # 预览内容随预览区大小自适应（图片按当前尺寸重新渲染）
        if self._preview_img is not None and self._preview_path:
            self._preview_image(self._preview_path)

    # ---------------- UI 辅助 ----------------
    def add_folder(self):
        idx = self._next_folder_id
        self._next_folder_id += 1
        block = FolderBlock(self.blocks_frame, self, idx)
        # treeview 行颜色需要在每个实例上单独配置
        block.tree.tag_configure("oddrow", background=COLORS["card"])
        block.tree.tag_configure("evenrow", background=COLORS["row_alt"])
        self.folder_blocks.append(block)
        # 稍后更新滚动区域
        self.root.after(50, lambda: self.canvas.configure(scrollregion=self.canvas.bbox("all")))

    def remove_folder(self, block):
        if block in self.folder_blocks:
            self.folder_blocks.remove(block)
        self.refresh_overall_counts()
        self.root.after(50, lambda: self.canvas.configure(scrollregion=self.canvas.bbox("all")))

    def _bind_mousewheel(self, widget):
        # 仅当内容高于 canvas 时滚动，避免空白区域时误滚动
        def on_mousewheel(event):
            bbox = widget.bbox("all")
            if bbox is None:
                return
            canvas_h = widget.winfo_height()
            content_h = bbox[3] - bbox[1]
            if content_h <= canvas_h:
                return
            if os.name == "nt":  # Windows 的 delta 是 120 的倍数
                widget.yview_scroll(-1 * (event.delta // 120), "units")
            else:  # macOS 等平台 delta 单位不同
                widget.yview_scroll(-1 * event.delta, "units")

        widget.bind_all("<MouseWheel>", on_mousewheel)
        if os.name == "posix":  # X11 系统使用 Button-4/Button-5
            widget.bind_all("<Button-4>", lambda e: widget.yview_scroll(-1, "units"))
            widget.bind_all("<Button-5>", lambda e: widget.yview_scroll(1, "units"))

    def refresh_overall_counts(self):
        total = sum(len(b.filtered_files) for b in self.folder_blocks if b.folder_path)
        # 当前已处理为 0（进度条在处理时显示）
        self.lbl_count.config(text=f"已处理 0 / {total}")

    # ---------------- 取消处理 ----------------
    def cancel_process(self):
        self.cancel_flag = True

    # ---------------- 文件预览 ----------------
    def show_preview(self, path):
        """根据文件路径刷新底部预览区；path 为空时显示占位提示。"""
        self._preview_path = path
        for child in self.preview_body.winfo_children():
            child.destroy()
        self._preview_img = None
        self._preview_rendered_size = None

        if not path or not os.path.isfile(path):
            self.lbl_preview_title.config(text="文件预览")
            self.lbl_preview_info.config(text="点击左侧文件列表中的文件查看预览")
            self._show_preview_placeholder("在文件列表中点击文件，即可在这里查看预览")
            return

        name = os.path.basename(path)
        ext = os.path.splitext(name)[1].lower()
        self.lbl_preview_title.config(text=name[:80])
        try:
            size = os.path.getsize(path)
        except OSError:
            size = 0
        self.lbl_preview_info.config(text=self._format_size(size))

        if ext in IMAGE_EXTS:
            self._preview_image(path)
        elif ext in TEXT_EXTS:
            self._preview_text(path)
        elif ext in TABLE_EXTS:
            self._preview_table(path, ext)
        else:
            # 无法识别的文件类型，尝试以文本模式预览
            message = f"暂不支持预览 {ext} 文件" if ext else "暂不支持预览该文件"
            self._try_text_fallback(path, message)

    def _show_preview_placeholder(self, message):
        tk.Label(
            self.preview_body, text=message, bg=COLORS["card"], fg=COLORS["text_secondary"],
            font=(FONT_FAMILY, 10),
        ).pack(expand=True)

    def _preview_image(self, path):
        try:
            from PIL import Image, ImageTk
        except ImportError:
            self._show_preview_placeholder("未安装 Pillow，无法预览图片\npip install pillow")
            return
        self.root.update_idletasks()
        max_w = max(self.preview_body.winfo_width() - 24, 100)
        max_h = max(self.preview_body.winfo_height() - 16, 60)
        if self._preview_rendered_size == (max_w, max_h):
            return  # 尺寸未变化，避免重复渲染
        self._preview_rendered_size = (max_w, max_h)
        try:
            img = Image.open(path)
            img.thumbnail((max_w, max_h))
            self._preview_img = ImageTk.PhotoImage(img)
        except Exception:
            self._preview_img = None
            self._preview_rendered_size = None
            self._try_text_fallback(path, "图片加载失败")
            return
        for child in self.preview_body.winfo_children():
            child.destroy()
        tk.Label(self.preview_body, image=self._preview_img, bg=COLORS["card"]).pack(padx=8, pady=6)

    def _preview_text(self, path):
        try:
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read(MAX_PREVIEW_BYTES)
        except Exception:
            self._try_text_fallback(path, "无法读取文本文件")
            return
        self._show_text_content(content)

    def _show_text_content(self, content):
        frame = tk.Frame(self.preview_body, bg=COLORS["card"])
        frame.pack(fill="both", expand=True, padx=8, pady=6)

        text = tk.Text(
            frame, wrap="none", bg=COLORS["preview_bg"], fg=COLORS["text"],
            insertbackground=COLORS["text"], font=MONO_FONT,
            relief="flat", borderwidth=0, padx=10, pady=8,
        )
        v_scroll = ttk.Scrollbar(frame, orient="vertical", command=text.yview)
        h_scroll = ttk.Scrollbar(frame, orient="horizontal", command=text.xview)
        text.configure(yscrollcommand=v_scroll.set, xscrollcommand=h_scroll.set)

        text.insert("1.0", content)
        if len(content) >= MAX_PREVIEW_BYTES:
            text.insert("end", "\n\n… 文件较大，仅预览前 200KB")
        text.configure(state="disabled")

        v_scroll.pack(side="right", fill="y")
        h_scroll.pack(side="bottom", fill="x")
        text.pack(side="left", fill="both", expand=True)

    def _try_text_fallback(self, path, fail_message):
        """无法按原类型预览时，尝试把文件当作文本读取。"""
        try:
            with open(path, "rb") as f:
                data = f.read(MAX_PREVIEW_BYTES)
        except Exception:
            self._show_preview_placeholder(fail_message)
            return
        if self._looks_binary(data):
            self._show_preview_placeholder(fail_message)
            return
        text = data.decode("utf-8", errors="replace")
        current = self.lbl_preview_info.cget("text")
        self.lbl_preview_info.config(text=f"{current} · 文本模式")
        self._show_text_content(text)

    @staticmethod
    def _looks_binary(data):
        if not data:
            return False
        sample = data[:8000]
        if b"\x00" in sample:
            return True
        control = sum(1 for b in sample if b < 32 and b not in (9, 10, 13, 12, 27))
        return control / len(sample) > 0.30

    def _preview_table(self, path, ext):
        if ext == ".xlsx":
            rows = self._read_xlsx(path)
        else:
            rows = []
            try:
                with open(path, "r", encoding="utf-8-sig", errors="replace", newline="") as f:
                    dialect = "excel-tab" if ext == ".tsv" else "excel"
                    for row in csv.reader(f, dialect):
                        rows.append(row)
                        if len(rows) >= MAX_PREVIEW_ROWS:
                            break
            except Exception:
                self._try_text_fallback(path, "无法读取表格文件")
                return
        if rows is None:
            return  # 读取失败时已显示提示
        self._show_table(rows)

    def _read_xlsx(self, path):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self._show_preview_placeholder("未安装 openpyxl，无法预览 Excel 表格\npip install openpyxl")
            return None
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
            ws = wb[wb.sheetnames[0]]
            sheet = ws.title
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append([self._cell_text(v) for v in row])
                if len(rows) >= MAX_PREVIEW_ROWS:
                    break
            wb.close()
            current = self.lbl_preview_info.cget("text")
            self.lbl_preview_info.config(text=f"{current} · 工作表: {sheet}")
            return rows
        except Exception:
            self._try_text_fallback(path, "Excel 文件读取失败")
            return None

    def _show_table(self, rows):
        if not rows:
            self._show_preview_placeholder("表格内容为空")
            return

        cols = min(max(len(r) for r in rows), 12)
        frame = tk.Frame(self.preview_body, bg=COLORS["card"])
        frame.pack(fill="both", expand=True, padx=8, pady=6)

        tree = ttk.Treeview(frame, columns=[f"c{i}" for i in range(cols)], show="headings", height=8)
        tree.tag_configure("odd", background=COLORS["card"])
        tree.tag_configure("even", background=COLORS["row_alt"])
        for i in range(cols):
            tree.heading(f"c{i}", text=f"列 {i + 1}")
            tree.column(f"c{i}", width=130, anchor="w", stretch=True)

        for idx, row in enumerate(rows):
            values = [""] * cols
            for i, v in enumerate(row[:cols]):
                values[i] = self._cell_text(v)
            tree.insert("", "end", values=values, tags=("odd" if idx % 2 == 0 else "even"))

        v_scroll = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        h_scroll = ttk.Scrollbar(frame, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=v_scroll.set, xscrollcommand=h_scroll.set)

        tree.grid(row=0, column=0, sticky="nsew")
        v_scroll.grid(row=0, column=1, sticky="ns")
        h_scroll.grid(row=1, column=0, sticky="ew")
        frame.rowconfigure(0, weight=1)
        frame.columnconfigure(0, weight=1)

    @staticmethod
    def _cell_text(value):
        if value is None:
            return ""
        return str(value).replace("\n", " ")[:60]

    @staticmethod
    def _format_size(size):
        if size < 1024:
            return f"{size} B"
        if size < 1024 * 1024:
            return f"{size / 1024:.1f} KB"
        return f"{size / 1024 / 1024:.1f} MB"

    # ---------------- 核心处理 ----------------
    def rename_and_save(self):
        prefix = self.entry_prefix.get().strip()
        if not prefix:
            messagebox.showwarning("警告", "请输入文件名前缀")
            return
        try:
            start_num = int(self.entry_start.get())
        except ValueError:
            messagebox.showwarning("警告", "起始数字必须为整数")
            return

        # 收集所有块中的文件
        items = []
        for b in self.folder_blocks:
            if b.folder_path:
                for name in b.filtered_files:
                    items.append((b.folder_path, name))
        if not items:
            messagebox.showwarning("警告", "没有可处理的文件")
            return

        save_folder = filedialog.askdirectory(title="选择保存目标文件夹")
        if not save_folder:
            return

        norm_save = os.path.normcase(os.path.normpath(save_folder))

        def is_inplace(folder):
            return os.path.normcase(os.path.normpath(folder)) == norm_save

        # 原位重命名模式下跳过“清空/新增”选择，避免清空源文件夹
        if not any(is_inplace(folder) for folder, _ in items):
            choice = messagebox.askyesnocancel(
                "保存模式选择",
                "请选择目标文件夹操作方式：\n\n是 = 清空文件夹\n否 = 在文件夹内新增\n取消 = 终止操作",
            )
            if choice is None:
                return
            if choice:
                try:
                    clear_files_in_folder(save_folder)
                except Exception as e:
                    messagebox.showerror("错误", f"清空目标文件夹失败: {e}")
                    return

        # 重置标志
        self.cancel_flag = False
        self.global_action = None

        # 预生成全部任务（编号与文件顺序一致）
        tasks = []  # (src, dst, inplace, src_name)
        cur_num = start_num
        for folder, name in items:
            ext = os.path.splitext(name)[1]
            new_name = f"{prefix}{cur_num}{ext}"
            inplace = is_inplace(folder)
            src = os.path.join(folder, name)
            dst = os.path.join(folder if inplace else save_folder, new_name)
            tasks.append((src, dst, inplace, name))
            cur_num += 1

        total = len(tasks)
        processed = 0
        moved = {}

        # 初始化进度条 UI
        self.progress["maximum"] = 100
        self.progress["value"] = 0
        self.lbl_percent.config(text="0%")
        self.lbl_count.config(text=f"已处理 0 / {total}")
        self.root.update_idletasks()

        # 原位重命名：先把目标名与其他待处理源文件冲突的文件挪到临时名
        inplace_tasks = [t for t in tasks if t[2]]
        if inplace_tasks:
            folder = os.path.dirname(inplace_tasks[0][0])
            inplace_jobs = [(t[3], os.path.basename(t[1])) for t in inplace_tasks]
            try:
                moved = prepare_inplace_temps(folder, inplace_jobs)
            except Exception as e:
                messagebox.showerror("错误", f"准备重命名失败: {e}")
                return
        else:
            folder = save_folder

        # 逐项处理
        try:
            for src, dst, inplace, src_name in tasks:
                if self.cancel_flag:
                    restore_inplace_temps(folder, moved)
                    messagebox.showinfo("已取消", "操作已取消")
                    self._update_progress(processed, total)
                    return

                # 原位重命名：文件已是指定名称则跳过
                if inplace and os.path.normcase(os.path.normpath(src)) == os.path.normcase(os.path.normpath(dst)):
                    processed += 1
                    self._update_progress(processed, total)
                    continue

                cur_src = moved.get(src_name) if inplace else None
                if cur_src is None:
                    cur_src = src

                # 处理冲突
                if os.path.exists(dst):
                    action = self._conflict_dialog(os.path.basename(dst))
                    if action == "cancel":
                        restore_inplace_temp(folder, src_name, cur_src)
                        restore_inplace_temps(folder, moved)
                        messagebox.showinfo("已取消", "操作已取消")
                        self._update_progress(processed, total)
                        return
                    if action == "skip":
                        # 若文件已被挪到临时名，则还原为原名
                        restore_inplace_temp(folder, src_name, cur_src)
                        moved.pop(src_name, None)
                        processed += 1
                        self._update_progress(processed, total)
                        continue
                    # 覆盖 => 继续处理

                if inplace:
                    os.replace(cur_src, dst)
                    moved.pop(src_name, None)
                else:
                    shutil.copy2(cur_src, dst)

                processed += 1
                self._update_progress(processed, total)
        except Exception as e:
            restore_inplace_temps(folder, moved)
            messagebox.showerror("错误", f"处理失败: {src}\n{e}")
            self._update_progress(processed, total)
            return

        # 处理完成
        if messagebox.askyesno("完成", "文件处理完成，是否现在打开目标文件夹？"):
            try:
                open_folder(save_folder)
            except Exception:
                pass

    def _conflict_dialog(self, filename):
        """返回 'overwrite'、'skip' 或 'cancel'。支持全局应用。"""
        if self.global_action == "overwrite":
            return "overwrite"
        if self.global_action == "skip":
            return "skip"

        choice = messagebox.askyesnocancel("文件已存在", f"{filename} 已存在。\n\n是 = 覆盖\n否 = 跳过\n取消 = 终止操作")
        if choice is None:
            return "cancel"
        elif choice:
            apply_all = messagebox.askyesno("应用到全部", "是否对后续冲突文件全部覆盖？")
            if apply_all:
                self.global_action = "overwrite"
            return "overwrite"
        else:
            apply_all = messagebox.askyesno("应用到全部", "是否对后续冲突文件全部跳过？")
            if apply_all:
                self.global_action = "skip"
            return "skip"

    def _update_progress(self, processed, total):
        percent = int((processed / total) * 100) if total else 0
        self.progress["value"] = percent
        self.lbl_percent.config(text=f"{percent}%")
        self.lbl_count.config(text=f"已处理 {processed} / {total}")
        self.root.update_idletasks()
